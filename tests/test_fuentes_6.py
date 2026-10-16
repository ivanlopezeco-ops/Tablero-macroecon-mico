"""Latencia de cada fuente: último dato disponible ahora y fuentes candidatas para datos del día."""
import datetime as dt, io, json, pathlib, traceback
import requests, urllib3
urllib3.disable_warnings()
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"}
out = [f"# Latencia de fuentes — {dt.datetime.utcnow():%Y-%m-%d %H:%M} UTC\n"]


def caso(nombre, f):
    out.append(f"\n## {nombre}")
    try:
        out.append(str(f())[:3000])
    except Exception:  # noqa: BLE001
        out.append("ERROR " + traceback.format_exc()[-800:])


hoy = dt.date.today()


def bcra_ids():
    res = []
    for i in [1, 78, 5, 1243, 15, 44, 108, 1266, 1187]:
        j = requests.get(f"https://api.bcra.gob.ar/estadisticas/v4.0/monetarias/{i}?desde={hoy-dt.timedelta(days=12)}&hasta={hoy}",
                         headers=UA, timeout=60, verify=False).json()
        det = j["results"][0]["detalle"]
        res.append(f"id {i}: " + ", ".join(f"{d['fecha']}={d['valor']}" for d in det[:4]))
    return "\n".join(res)


caso("BCRA API v4 (últimos datos)", bcra_ids)


def series_xlsm():
    import openpyxl
    r = requests.get("https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/series.xlsm", headers=UA, timeout=120, verify=False)
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    res = [f"tamaño {len(r.content)} | hojas: {wb.sheetnames}"]
    for h in wb.sheetnames:
        if any(k in h.upper() for k in ["RESERV", "INSTRU", "BASE"]):
            filas = [x for x in wb[h].iter_rows(values_only=True) if any(v is not None for v in x[:6])]
            res.append(f"--- {h}: últimas filas")
            for x in filas[-4:]:
                res.append(repr(x[:12]))
    return "\n".join(res)


caso("BCRA series.xlsm (planilla diaria)", series_xlsm)
caso("ArgentinaDatos riesgo país último", lambda: requests.get(
    "https://api.argentinadatos.com/v1/finanzas/indices/riesgo-pais/ultimo", headers=UA, timeout=30).text)
caso("ArgentinaDatos riesgo país serie (cola)", lambda: requests.get(
    "https://api.argentinadatos.com/v1/finanzas/indices/riesgo-pais", headers=UA, timeout=30).json()[-4:])
caso("Ámbito riesgo país variación", lambda: requests.get(
    "https://mercados.ambito.com//riesgopais/variacion", headers=UA, timeout=30).text)
caso("Ámbito riesgo país histórico", lambda: requests.get(
    f"https://mercados.ambito.com//riesgopais/historico-general/{hoy-dt.timedelta(days=7)}/{hoy}", headers=UA, timeout=30).text)

BYMA = "https://open.bymadata.com.ar/vanoms-be-core/rest/api/bymadata/free"
BH = {**UA, "Content-Type": "application/json", "Origin": "https://open.bymadata.com.ar", "Referer": "https://open.bymadata.com.ar/"}


def byma_bonos():
    r = requests.post(f"{BYMA}/public-bonds", json={"excludeZeroPxAndQty": True, "T2": False, "T1": True, "T0": False,
                      "Content-Type": "application/json"}, headers=BH, verify=False, timeout=60)
    j = r.json()
    data = j.get("data", j) if isinstance(j, dict) else j
    sel = [d for d in data if str(d.get("symbol", "")).startswith(("AL30", "GD30", "GD35", "AE38"))]
    return f"status {r.status_code} n={len(data)}\n" + "\n".join(json.dumps(d)[:600] for d in sel[:8])


caso("BYMA bonos públicos (precio intradiario)", byma_bonos)


def byma_hist_hoy():
    t0 = int(dt.datetime.combine(hoy - dt.timedelta(days=5), dt.time()).timestamp())
    j = requests.get(f"{BYMA}/chart/historical-series/history", verify=False, headers=UA, timeout=60,
                     params={"symbol": "AL30D 24HS", "resolution": "D", "from": t0, "to": int(dt.datetime.now().timestamp())}).json()
    return list(zip(j.get("t", []), j.get("c", [])))


caso("BYMA histórico AL30D (¿incluye hoy?)", byma_hist_hoy)


def mae_forex():
    p = {"oTitulo": json.dumps({"fechaDesde": str(hoy), "fechaHasta": str(hoy)})}
    H = {**UA, "Origin": "https://marketdata.mae.com.ar", "Referer": "https://marketdata.mae.com.ar/"}
    a = requests.get("https://api.marketdata.mae.com.ar/api/mercado/titulo/historicoforex/volumenoperado", headers=H, params=p, timeout=60).text[:800]
    b = requests.get("https://api.marketdata.mae.com.ar/api/mercado/forex", headers=H, timeout=60).text[:1500]
    return a + "\n--- forex en vivo:\n" + b


caso("A3/MAE forex hoy", mae_forex)

pathlib.Path("results").mkdir(exist_ok=True)
pathlib.Path("results/test_report_6.md").write_text("\n".join(out), encoding="utf-8")
